from openpyxl import workbook, load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
import sqlalchemy
from sqlalchemy import text
import string
from openpyxl.workbook.defined_name import DefinedName

def generar_formulario_excel(file_path, save_file_path, id_empresa):
    def gen_col_letter():
        letter = dict(zip(range(1, 27), string.ascii_uppercase))
        alphabet = string.ascii_uppercase
        letters = []
        numbers = []
        for i in range(0,26*27):
            x = i//26
            if x == 0:
                character = alphabet[i]
            else:
                y = i-(26*x)
                character = letter[x] + alphabet[y]
            letters.append(character)
            numbers.append(i)

        res = {numbers[i]+1: letters[i] for i in range(len(letters))}
        return res


    col_letter = gen_col_letter()

    # Acceder a la base de datos
    username = 'lguareschi'
    password = 'L01SGSc1r*20'
    ip = '172.16.24.35'
    port = '3306'
    db_name = 'ubiimarket_db'
    db = sqlalchemy.create_engine('mysql+pymysql://{}:{}@{}:{}/{}'.format(username, password, ip, port, db_name),
                                  pool_recycle=3600)

    # Hacer query en la bd
    query = text('SELECT * '
                 'FROM ubiimarket_db.dt_empresa_ramo as a, ubiimarket_db.tm_ramo as b '
                 'WHERE a.id_ramo=b.id_ramo and a.id_empresa={};'.format(id_empresa))
    ramos = db.engine.execute(query).fetchall()

    # Abrir el archivo excel
    # Start by opening the spreadsheet and selecting the main sheet
    excel_file = load_workbook(filename=file_path)
    sheet = excel_file['LISTS']

    # Agregar Ramos
    row_ramo = 1
    for ramo in ramos:
        row_ramo += 1
        cell = 'A' + str(row_ramo)
        sheet[cell] = ramo['strnombre_ramo']
    n_ramos = row_ramo - 1

    # Agregar titulos de categorias
    col = 1
    for ramo in ramos:
        col += 1
        # Crear titulo
        cell = col_letter[col] + '1'
        table_end_cell = col_letter[col] + '2'
        sheet[cell] = ramo['strnombre_ramo']

        # Crear titulo de la tabla
        table_name = 'T_' + ramo['strnombre_ramo']
        table_name = table_name.replace(' ', '_')
        table = Table(displayName=table_name, ref=cell + ':' + table_end_cell)
        style = TableStyleInfo(name="TableStyleLight1", showFirstColumn=False,
                               showLastColumn=False, showRowStripes=True, showColumnStripes=False)
        table.tableStyleInfo = style
        sheet.add_table(table)

        # Crear nombre de rango de valores
        range_name = ramo['strnombre_ramo'].replace(' ', '_')
        new_range = DefinedName(range_name, attr_text='{}[{}]'.format(table_name, ramo['strnombre_ramo']))
        excel_file.defined_names.append(new_range)

    # Agregar categoria de cada ramo
    col = 1
    for ramo in ramos:
        col += 1
        row_categoria = 1
        query_categoria = text('SELECT cat.id_ramo, ramo.strnombre_ramo, cat.strnombre_categoria '
                               'FROM ubiimarket_db.tm_categoria as cat, ubiimarket_db.tm_ramo as ramo '
                               'WHERE ramo.id_ramo=cat.id_ramo AND strnombre_ramo=\'{}\''.format(ramo['strnombre_ramo']))
        categorias = db.engine.execute(query_categoria).fetchall()
        for categoria in categorias:
            row_categoria += 1
            cell = col_letter[col] + str(row_categoria)
            sheet[cell] = categoria['strnombre_categoria']

    # Agregar marcas
    query_marcas = text('SELECT * '
                        'FROM ubiimarket_db.tm_marca as a, ubiimarket_db.tm_ramo as b '
                        'WHERE a.id_ramo=b.id_ramo and id_empresa={}'.format(id_empresa))
    marcas = db.engine.execute(query_marcas).fetchall()
    row = 1
    col += 1
    added_marcas = []
    cell = col_letter[col] + str(row)
    if row == 1:
        sheet[cell] = 'MARCAS'
    # Crear titulo de la tabla marcas
    table_name = 'T_MARCAS'
    table = Table(displayName=table_name, ref=col_letter[col] + '1:' + col_letter[col] + '2')
    style = TableStyleInfo(name="TableStyleLight1", showFirstColumn=False,
                           showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    table.tableStyleInfo = style
    sheet.add_table(table)

    # Crear nombre de rango de valores
    range_name = 'MARCAS'
    new_range = DefinedName(range_name, attr_text='{}[{}]'.format(table_name, range_name))
    excel_file.defined_names.append(new_range)

    for marca in marcas:
        row += 1
        cell = col_letter[col] + str(row)
        nombre_marca = marca['strnombre_marca']
        if row > 1 and nombre_marca not in added_marcas:
            sheet[cell] = nombre_marca
            added_marcas.append(nombre_marca)
        elif row > 1 and nombre_marca in added_marcas:
            row = row - 1

    # Crear modelos de cada marca
    added_marcas = []
    for marca in marcas:
        row = 1
        col += 1
        id_marca = marca['id_marca']
        nombre_marca = marca['strnombre_marca']
        query_modelos = text('SELECT * FROM ubiimarket_db.tm_modelo WHERE id_marca={}'.format(id_marca))
        modelos = db.engine.execute(query_modelos).fetchall()
        if row == 1 and nombre_marca not in added_marcas:
            cell = col_letter[col] + str(row)
            sheet[cell] = nombre_marca
            # Crear titulo de la tabla marcas
            table_name = 'T_' + nombre_marca
            table_name = table_name.replace(' ', '_')
            table = Table(displayName=table_name, ref=col_letter[col] + '1:' + col_letter[col] + '2')
            style = TableStyleInfo(name="TableStyleLight1", showFirstColumn=False,
                                   showLastColumn=False, showRowStripes=True, showColumnStripes=False)
            table.tableStyleInfo = style
            sheet.add_table(table)
            added_marcas.append(nombre_marca)
            add_modelo = True
            # Crear nombre de rango de valores
            range_name = nombre_marca.replace(' ', '_')
            new_range = DefinedName(range_name, attr_text='{}[{}]'.format(table_name, nombre_marca))
            excel_file.defined_names.append(new_range)
        else:
            add_modelo = False
            col = col - 1
        if add_modelo:
            for modelo in modelos:
                row += 1
                cell = col_letter[col] + str(row)
                nombre_modelo = modelo['strnombre_modelo']
                if row > 1:
                    sheet[cell] = nombre_modelo

    # Agregar IVAs
    query_ivas = text('SELECT * FROM ubiimarket_db.tm_iva;')
    ivas = db.engine.execute(query_ivas).fetchall()
    col += 1
    row = 1
    cell = col_letter[col] + str(row)
    # Agregar titulo de celda
    sheet[cell] = 'IVAS'
    # Agregar tabla
    table_name = 'T_' + 'IVAS'
    table = Table(displayName=table_name, ref=col_letter[col] + '1:' + col_letter[col] + '2')
    style = TableStyleInfo(name="TableStyleLight1", showFirstColumn=False,
                           showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    table.tableStyleInfo = style
    sheet.add_table(table)
    # Crear nombre de rango de valores
    range_name = 'IVAS'
    new_range = DefinedName(range_name, attr_text='{}[{}]'.format(table_name, 'IVAS'))
    excel_file.defined_names.append(new_range)
    for iva in ivas:
        row += 1
        cell = col_letter[col] + str(row)
        numero_iva = iva['porcentaje_iva']
        sheet[cell] = numero_iva

    # Agregar unidad de despacho
    query_und_despacho = text('SELECT * FROM ubiimarket_db.tm_unidad_despacho;')
    unidades_despacho = db.engine.execute(query_und_despacho).fetchall()
    col += 1
    row = 1
    cell = col_letter[col] + str(row)
    # Agregar titulo de celda
    sheet[cell] = 'UNIDAD DESPACHO'
    table_name = 'T_' + 'UNIDAD_DESPACHO'
    table = Table(displayName=table_name, ref=col_letter[col] + '1:' + col_letter[col] + '2')
    style = TableStyleInfo(name="TableStyleLight1", showFirstColumn=False,
                           showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    table.tableStyleInfo = style
    sheet.add_table(table)
    # Crear nombre de rango de valores
    range_name = 'UNIDAD_DESPACHO'
    new_range = DefinedName(range_name, attr_text='{}[{}]'.format(table_name, 'UNIDAD DESPACHO'))
    excel_file.defined_names.append(new_range)
    for unidad in unidades_despacho:
        row += 1
        cell = col_letter[col] + str(row)
        nombre_unidad = unidad['strnombre_despacho'].upper()
        sheet[cell] = nombre_unidad

    # Agregar unidad de presentacion
    query_und_presentacion = text('SELECT * FROM ubiimarket_db.tm_unidad_presentacion;')
    unidades_presentacion = db.engine.execute(query_und_presentacion).fetchall()
    col += 1
    row = 1
    cell = col_letter[col] + str(row)
    # Agregar titulo de celda
    sheet[cell] = 'UNIDAD PRESENTACION'
    table_name = 'T_' + 'UNIDAD_PRESENTACION'
    table = Table(displayName=table_name, ref=col_letter[col] + '1:' + col_letter[col] + '2')
    style = TableStyleInfo(name="TableStyleLight1", showFirstColumn=False,
                           showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    table.tableStyleInfo = style
    sheet.add_table(table)
    # Crear nombre de rango de valores
    range_name = 'UNIDAD_PRESENTACION'
    new_range = DefinedName(range_name, attr_text='{}[{}]'.format(table_name, 'UNIDAD PRESENTACION'))
    excel_file.defined_names.append(new_range)
    for unidad in unidades_presentacion:
        row += 1
        cell = col_letter[col] + str(row)
        nombre_unidad = unidad['strnombre_presentacion'].upper()
        sheet[cell] = nombre_unidad

    # Agregar SI/NO
    col += 1
    row = 1
    cell = col_letter[col] + str(row)
    # Agregar titulo de celda
    sheet[cell] = 'SI NO'
    table_name = 'T_' + 'SI_NO'
    table = Table(displayName=table_name, ref=col_letter[col] + '1:' + col_letter[col] + '2')
    style = TableStyleInfo(name="TableStyleLight1", showFirstColumn=False,
                           showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    table.tableStyleInfo = style
    sheet.add_table(table)
    # Crear nombre de rango de valores
    range_name = 'SI_NO'
    new_range = DefinedName(range_name, attr_text='{}[{}]'.format(table_name, 'SI NO'))
    excel_file.defined_names.append(new_range)
    # Agregar filas
    for i in range(0,2):
        row += 1
        cell = col_letter[col] + str(row)
        if i == 0:
            sheet[cell] = 'SI'
        elif i == 1:
            sheet[cell] = 'NO'

    # Agregar codigos de productos
    query_cod_productos = text('SELECT * FROM ubiimarket_db.dt_productos WHERE id_empresa={};'.format(id_empresa))
    codigos_prodcutos = db.engine.execute(query_cod_productos).fetchall()
    col += 1
    row = 1
    # Agregar titulo de celda
    cell = col_letter[col] + str(row)
    sheet[cell] = 'CODIGOS PRODUCTOS'
    table_name = 'T_' + 'CODIGOS_PRODUCTOS'
    table = Table(displayName=table_name, ref=col_letter[col] + '1:' + col_letter[col] + '2')
    style = TableStyleInfo(name="TableStyleLight1", showFirstColumn=False,
                           showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    table.tableStyleInfo = style
    sheet.add_table(table)
    # Crear nombre de rango de valores
    range_name = 'CODIGOS_PRODUCTOS'
    new_range = DefinedName(range_name, attr_text='{}[{}]'.format(table_name, 'CODIGOS PRODUCTOS'))
    excel_file.defined_names.append(new_range)
    for codigo in codigos_prodcutos:
        row += 1
        cell = col_letter[col] + str(row)
        n_codigo = codigo['CodProd']
        sheet[cell] = n_codigo

    # Agregar RIF
    query_rif = text('SELECT * FROM ubiimarket_db.dt_empresa WHERE id_empresa={};'.format(id_empresa))
    rifs = db.engine.execute(query_rif).fetchall()
    col += 1
    row = 1
    # Agregar titulo de celda
    cell = col_letter[col] + str(row)
    sheet[cell] = 'RIF'
    table_name = 'T_' + 'RIF'
    table = Table(displayName=table_name, ref=col_letter[col] + '1:' + col_letter[col] + '2')
    style = TableStyleInfo(name="TableStyleLight1", showFirstColumn=False,
                           showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    table.tableStyleInfo = style
    sheet.add_table(table)
    # Crear nombre de rango de valores
    range_name = 'RIF'
    new_range = DefinedName(range_name, attr_text='{}[{}]'.format(table_name, 'RIF'))
    excel_file.defined_names.append(new_range)
    for rif in rifs:
        row += 1
        cell = col_letter[col] + str(row)
        n_rif = rif['strrif_empresa']
        sheet[cell] = n_rif

    # Formatear todas las tablas
    col = 0
    for table in sheet.tables.values():
        col += 1
        n_rows = 0
        for cell in sheet[col_letter[col]]:
            if cell.value is not None:
                n_rows += 1
        if n_rows < 2:
            n_rows = 2
        table.ref = "{}1:{}{}".format(col_letter[col], col_letter[col], n_rows)

    # Guardar el archivo excel (ULTIMO PASO)
    excel_file.save(filename=save_file_path)

'''==============CODIGO DE PRUEBA=============='''
file_path = 'C:\\Users\\luisg\\Documents\\Python Projects\\Git_Ubii Company\\Creador Excel Carga_masiva\\CARGAMASIVA_V3.0.xlsx'
save_file_path = 'output.xlsx'
# id de la empresa, para mostrar los datos dependiendo de cada usuario
id_empresa = 157

generar_formulario_excel(file_path, save_file_path, id_empresa)
print('TABLA CREADA EXITOSAMENTE')

'''
==============================================================================
PARA VER LAS SHEETS OCULTAS DE EXCEL:
-Seleccionar vista de codigo y desocultarlas de ahi directamente
==============================================================================
'''

